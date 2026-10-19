
from io import BytesIO
from datetime import datetime

from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .forms import IncidenciaForm
from .models import Incidencia, PerfilUsuario


def obtener_perfil(usuario):
    try:
        return usuario.perfil
    except PerfilUsuario.DoesNotExist:
        raise PermissionDenied(
            "Tu usuario no tiene un perfil configurado. "
            "Contacta al administrador de Sistemas."
        )


@login_required
def inicio(request):
    perfil = obtener_perfil(request.user)

    if perfil.rol == PerfilUsuario.Roles.ADMIN:
        incidencias = Incidencia.objects.select_related(
            "sucursal", "creado_por", "responsable"
        ).all()
    else:
        if perfil.sucursal is None:
            raise PermissionDenied(
                "Tu usuario no tiene una sucursal asignada."
            )

        incidencias = Incidencia.objects.filter(
            sucursal=perfil.sucursal
        ).select_related(
            "sucursal", "creado_por", "responsable"
        )

    return render(
        request,
        "incidencias/inicio.html",
        {
            "incidencias": incidencias,
            "perfil": perfil,
        },
    )


@login_required
def nueva_incidencia(request):
    perfil = obtener_perfil(request.user)

    if perfil.rol != PerfilUsuario.Roles.ADMIN:
        if perfil.sucursal is None:
            raise PermissionDenied(
                "Tu usuario no tiene una sucursal asignada."
            )

    if request.method == "POST":
        formulario = IncidenciaForm(request.POST)

        if formulario.is_valid():
            incidencia = formulario.save(commit=False)
            incidencia.creado_por = request.user

            if perfil.rol == PerfilUsuario.Roles.ADMIN:
                raise PermissionDenied(
                    "El registro desde esta pantalla está habilitado "
                    "para usuarios de sucursal."
                )

            incidencia.sucursal = perfil.sucursal
            incidencia.save()

            return redirect("inicio")
    else:
        formulario = IncidenciaForm()

    return render(
        request,
        "incidencias/nueva_incidencia.html",
        {
            "formulario": formulario,
            "perfil": perfil,
        },
    )


@login_required
def exportar_excel(request):
    perfil = obtener_perfil(request.user)

    if perfil.rol == PerfilUsuario.Roles.ADMIN:
        incidencias = Incidencia.objects.select_related(
            "sucursal", "creado_por", "responsable"
        ).all()
    else:
        if perfil.sucursal is None:
            raise PermissionDenied(
                "Tu usuario no tiene una sucursal asignada."
            )

        incidencias = Incidencia.objects.filter(
            sucursal=perfil.sucursal
        ).select_related(
            "sucursal", "creado_por", "responsable"
        )

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Pendientes"

    encabezados = [
        "Folio", "Sucursal", "Tipo", "Categoría", "Título",
        "Descripción", "Prioridad", "Estado", "Creado por",
        "Responsable", "Fecha compromiso", "Solución",
        "Fecha de registro", "Última actualización",
        "Fecha de resolución",
    ]
    hoja.append(encabezados)

    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(
            fill_type="solid", fgColor="1F4E78"
        )
        celda.alignment = Alignment(
            horizontal="center", vertical="center"
        )

    def formato_fecha(valor):
        if not valor:
            return ""

        if isinstance(valor, datetime):
            if timezone.is_aware(valor):
                valor = timezone.localtime(valor)
            return valor.strftime("%d/%m/%Y %H:%M")

        if hasattr(valor, "strftime"):
            return valor.strftime("%d/%m/%Y")

        return str(valor)

    for incidencia in incidencias.order_by("-creado_en"):
        responsable = incidencia.responsable
        creador = incidencia.creado_por

        hoja.append([
            incidencia.folio,
            str(incidencia.sucursal),
            incidencia.get_tipo_display(),
            incidencia.categoria,
            incidencia.titulo,
            incidencia.descripcion,
            incidencia.get_prioridad_display(),
            incidencia.get_estado_display(),
            (
                creador.get_full_name() or creador.username
                if creador else ""
            ),
            (
                responsable.get_full_name() or responsable.username
                if responsable else ""
            ),
            formato_fecha(incidencia.fecha_compromiso),
            incidencia.solucion or "",
            formato_fecha(incidencia.creado_en),
            formato_fecha(incidencia.actualizado_en),
            formato_fecha(incidencia.resuelto_en),
        ])

    for columna in hoja.columns:
        letra = columna[0].column_letter
        ancho = max(
            len(str(celda.value or ""))
            for celda in columna
        )
        hoja.column_dimensions[letra].width = min(ancho + 3, 45)

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    respuesta = HttpResponse(
        archivo.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    respuesta["Content-Disposition"] = (
        'attachment; filename="reporte_pendientes.xlsx"'
    )

    return respuesta
